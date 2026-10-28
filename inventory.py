# -*- coding: utf-8 -*-
"""
设备清单读取模块
支持直接读取 Excel 文件（.xlsx / .xls）作为设备清单
Excel 列说明（列名不区分大小写，顺序随意）：
  - IP地址 / ip          必填
  - 用户名 / username    必填
  - 密码   / password    必填
  - 主机名 / hostname    可选，为空时首次备份后自动回填
  - 端口   / port        可选，默认 22
  - 厂商   / vendor      可选，默认 auto
其他未识别的列会被忽略，不影响读取。
"""

import openpyxl
import re
import os
import shutil
from copy import copy
from typing import List, Dict, Optional


# Excel 列名到标准字段的映射（支持中英文、模糊匹配）
_COL_MAP = {
    "ip":         ["ip", "ip地址", "管理ip", "设备ip", "ip address"],
    "username":   ["username", "用户名", "账号", "user", "登录名"],
    "password":   ["password", "密码", "口令", "passwd", "pwd"],
    "hostname":   ["hostname", "主机名", "设备名", "名称", "host"],
    "port":       ["port", "端口", "ssh端口", "ssh_port"],
    "vendor":     ["vendor", "厂商", "品牌", "设备厂商"],
}


def _normalize(s: str) -> str:
    """列名标准化：小写 + 去空格"""
    return re.sub(r"\s+", "", str(s or "")).lower()


def _match_column(header: str) -> Optional[str]:
    """将 Excel 列名匹配到标准字段名"""
    h = _normalize(header)
    for field, aliases in _COL_MAP.items():
        if h in [_normalize(a) for a in aliases]:
            return field
    return None


def load_from_excel(path: str, sheet_name: str = None,
                    default_username: str = "admin",
                    default_password: str = "",
                    default_port: int = 22,
                    default_vendor: str = "auto") -> List[Dict]:
    """
    从 Excel 文件读取设备清单，返回与 TOML devices 列表相同格式的字典列表。

    参数：
        path            - Excel 文件路径
        sheet_name      - 指定 Sheet 名（默认读取第一个 Sheet）
        default_*       - 当 Excel 中无对应列时使用的默认值

    返回：
        [{"ip": ..., "username": ..., "password": ..., ...}, ...]
    """
    if not os.path.exists(path):
        raise FileNotFoundError(f"Excel 文件不存在: {path}")

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

    # 选 Sheet
    if sheet_name:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Sheet '{sheet_name}' 不存在，可用: {wb.sheetnames}")
        ws = wb[sheet_name]
    else:
        ws = wb.active
        print(f"  [Excel] 读取 Sheet: {ws.title}")

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("Excel 文件为空")

    # ---------- 识别表头行 ----------
    header_row_idx = None
    col_map: Dict[int, str] = {}   # {列索引: 标准字段名}

    for row_idx, row in enumerate(rows[:5]):   # 只在前5行找表头
        matched = {}
        for col_idx, cell in enumerate(row):
        	field = _match_column(str(cell or ""))
        	if field:
        		matched[col_idx] = field
        if "ip" in matched.values() and "password" in matched.values():
            header_row_idx = row_idx
            col_map = matched
            break

    if header_row_idx is None:
        raise ValueError(
            "未找到有效表头行！请确保 Excel 中包含 'IP地址'、'用户名'、'密码' 等列。"
        )

    print(f"  [Excel] 表头识别成功（第{header_row_idx+1}行）: "
          f"{list(col_map.values())}")

    # ---------- 读取数据行 ----------
    devices = []
    skipped = 0

    for row in rows[header_row_idx + 1:]:
        if all(cell is None or str(cell).strip() == "" for cell in row):
            continue   # 跳过全空行

        def get(field_name):
            for col_idx, f in col_map.items():
                if f == field_name and col_idx < len(row):
                    val = row[col_idx]
                    return str(val).strip() if val is not None else ""
            return ""

        ip = get("ip")
        if not ip or not re.match(r"\d+\.\d+\.\d+\.\d+", ip):
            skipped += 1
            continue   # 跳过无效IP行

        username = get("username") or default_username
        password = get("password") or default_password

        if not username or not password:
            print(f"  [警告] {ip} 缺少用户名或密码，跳过")
            skipped += 1
            continue

        port_str = get("port")
        try:
            port = int(port_str) if port_str else default_port
        except ValueError:
            port = default_port

        vendor = get("vendor") or default_vendor
        # 归一化厂商名
        vendor_lower = vendor.lower()
        if "华为" in vendor_lower or "huawei" in vendor_lower:
            vendor = "huawei"
        elif "h3c" in vendor_lower or "华三" in vendor_lower:
            vendor = "h3c"
        elif vendor_lower not in ("auto", "huawei", "h3c"):
            vendor = "auto"

        devices.append({
            "ip":       ip,
            "username": username,
            "password": password,
            "port":     port,
            "vendor":   vendor,
            "hostname": get("hostname"),
            "_orig_ip": ip,      # GUI 就地编辑时用它定位 Excel 行
        })

    wb.close()
    print(f"  [Excel] 读取完成：有效设备 {len(devices)} 台，跳过 {skipped} 行")
    return devices


def write_back_hostnames(path: str, hostnames: Dict[str, str],
                         sheet_name: str = None) -> Dict:
    """
    把设备主机名回填到 Excel 清单的“主机名”列。

    设计取舍：
      - 只填空单元格，绝不覆盖用户手工维护的主机名
      - 清单里没有“主机名”列时，自动在末尾追加一列
      - 先写同目录临时文件再 os.replace 原子替换，并保留一份 .bak
      - 文件被 Excel 占用时返回错误，不抛异常

    返回 {"filled": n, "skipped": m, "error": ""}
    """
    info = {"filled": 0, "skipped": 0, "error": ""}
    if not hostnames:
        return info
    if not os.path.exists(path):
        info["error"] = f"文件不存在: {path}"
        return info

    try:
        wb = openpyxl.load_workbook(path)
    except Exception as e:
        info["error"] = f"无法打开 Excel: {e}"
        return info

    try:
        if sheet_name:
            if sheet_name not in wb.sheetnames:
                info["error"] = f"Sheet '{sheet_name}' 不存在"
                return info
            ws = wb[sheet_name]
        else:
            ws = wb.active

        # 定位表头行（前 5 行内，必须有 IP 列）
        header_row, col_map = None, {}
        for row_idx, row in enumerate(
                ws.iter_rows(min_row=1, max_row=min(5, ws.max_row),
                             values_only=True), 1):
            matched = {}
            for col_idx, cell in enumerate(row, 1):
                field = _match_column(str(cell or ""))
                if field:
                    matched[col_idx] = field
            if "ip" in matched.values():
                header_row, col_map = row_idx, matched
                break
        if header_row is None:
            info["error"] = "未找到表头行（缺少 IP 列）"
            return info

        ip_col = next(c for c, f in col_map.items() if f == "ip")
        hn_col = next((c for c, f in col_map.items() if f == "hostname"), None)

        # 没有主机名列：在末尾补一列，沿用 IP 表头的样式
        if hn_col is None:
            hn_col = ws.max_column + 1
            cell = ws.cell(row=header_row, column=hn_col, value="主机名")
            ref = ws.cell(row=header_row, column=ip_col)
            cell.fill = copy(ref.fill)
            cell.font = copy(ref.font)
            cell.alignment = copy(ref.alignment)

        for row in range(header_row + 1, ws.max_row + 1):
            ip = str(ws.cell(row=row, column=ip_col).value or "").strip()
            name = str(hostnames.get(ip) or "").strip()
            if not name:
                continue
            cell = ws.cell(row=row, column=hn_col)
            if str(cell.value or "").strip():
                info["skipped"] += 1      # 已有手工维护的值，保留
                continue
            cell.value = name
            info["filled"] += 1

        if not info["filled"]:
            wb.close()
            return info

        try:
            shutil.copy2(path, path + ".bak")
        except Exception:
            pass

        tmp = path + ".tmp"
        try:
            wb.save(tmp)
        finally:
            wb.close()
        try:
            os.replace(tmp, path)
        except Exception:
            try:
                os.remove(tmp)
            except Exception:
                pass
            info["error"] = "清单文件被占用，请先在 Excel 中关闭后重新运行"
            info["filled"] = 0
        return info
    except Exception as e:
        try:
            wb.close()
        except Exception:
            pass
        info["error"] = f"回填失败: {e}"
        return info


def update_device_in_excel(path: str, orig_ip: str, updates: Dict,
                           sheet_name: str = None) -> Dict:
    """
    修改清单中某台设备的字段（ip / hostname / port），用于 GUI 就地编辑。
    按 orig_ip 定位行，只更新 updates 中提供的字段；缺少的列会自动追加。
    返回 {"updated": n, "error": ""}
    """
    info = {"updated": 0, "error": ""}
    if not os.path.exists(path):
        info["error"] = f"文件不存在: {path}"
        return info
    try:
        wb = openpyxl.load_workbook(path)
    except Exception as e:
        info["error"] = f"无法打开 Excel: {e}"
        return info

    try:
        if sheet_name:
            if sheet_name not in wb.sheetnames:
                info["error"] = f"Sheet '{sheet_name}' 不存在"
                return info
            ws = wb[sheet_name]
        else:
            ws = wb.active

        header_row, col_map = None, {}
        for row_idx, row in enumerate(
                ws.iter_rows(min_row=1, max_row=min(5, ws.max_row),
                             values_only=True), 1):
            matched = {}
            for col_idx, cell in enumerate(row, 1):
                field = _match_column(str(cell or ""))
                if field:
                    matched[col_idx] = field
            if "ip" in matched.values():
                header_row, col_map = row_idx, matched
                break
        if header_row is None:
            info["error"] = "未找到表头行（缺少 IP 列）"
            return info

        ip_col = next(c for c, f in col_map.items() if f == "ip")

        def ensure_col(field: str, header: str) -> int:
            col = next((c for c, f in col_map.items() if f == field), None)
            if col is None:
                col = ws.max_column + 1
                cell = ws.cell(row=header_row, column=col, value=header)
                ref = ws.cell(row=header_row, column=ip_col)
                cell.fill = copy(ref.fill)
                cell.font = copy(ref.font)
                cell.alignment = copy(ref.alignment)
                col_map[col] = field
            return col

        target_row = None
        for row in range(header_row + 1, ws.max_row + 1):
            if str(ws.cell(row=row, column=ip_col).value or "").strip() == orig_ip:
                target_row = row
                break
        if target_row is None:
            info["error"] = f"清单中找不到原 IP: {orig_ip}"
            return info

        field_cols = {}
        if "ip" in updates:
            field_cols["ip"] = ip_col
        if "hostname" in updates:
            field_cols["hostname"] = ensure_col("hostname", "主机名")
        if "port" in updates:
            field_cols["port"] = ensure_col("port", "端口")

        for field, col in field_cols.items():
            new_val = updates[field]
            if field == "port":
                try:
                    new_val = int(new_val)
                except Exception:
                    info["error"] = "端口必须是整数"
                    return info
                if not (1 <= new_val <= 65535):
                    info["error"] = "端口必须在 1-65535 之间"
                    return info
            cell = ws.cell(row=target_row, column=col)
            old_val = cell.value if cell.value is not None else ""
            if str(old_val) != str(new_val):
                cell.value = new_val
                info["updated"] += 1

        if not info["updated"]:
            wb.close()
            return info

        try:
            shutil.copy2(path, path + ".bak")
        except Exception:
            pass
        tmp = path + ".tmp"
        try:
            wb.save(tmp)
        finally:
            wb.close()
        try:
            os.replace(tmp, path)
        except Exception:
            try:
                os.remove(tmp)
            except Exception:
                pass
            info["error"] = "清单文件被占用，请先在 Excel 中关闭后重试"
            info["updated"] = 0
        return info
    except Exception as e:
        try:
            wb.close()
        except Exception:
            pass
        info["error"] = f"保存失败: {e}"
        return info


def create_excel_template(output_path: str):
    """
    生成一个标准的设备清单 Excel 模板，方便用户直接填写或与现有表格对照。
    """
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "设备清单"

    headers = ["IP地址", "用户名", "密码", "主机名", "端口"]
    examples = [
        ["192.168.1.1", "admin", "Admin@123", "", 22],
        ["192.168.1.2", "admin", "Admin@123", "", 22],
        ["192.168.2.1", "admin", "Admin@456", "", 22],
    ]

    fill_header = PatternFill("solid", fgColor="1F4E79")
    fill_example = PatternFill("solid", fgColor="EBF5FB")
    font_header = Font(bold=True, color="FFFFFF", size=11)
    align_center = Alignment(horizontal="center", vertical="center")

    ws.row_dimensions[1].height = 22
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.fill = fill_header
        c.font = font_header
        c.alignment = align_center

    for row_idx, row_data in enumerate(examples, 2):
        for col_idx, val in enumerate(row_data, 1):
            c = ws.cell(row=row_idx, column=col_idx, value=val)
            c.fill = fill_example
            c.alignment = Alignment(vertical="center")

    # 自动列宽
    col_widths = [16, 12, 14, 22, 8]
    for col, w in enumerate(col_widths, 1):
        ws.column_dimensions[ws.cell(1, col).column_letter].width = w

    # 添加说明 Sheet
    ws2 = wb.create_sheet("填写说明")
    notes = [
        ("列名", "说明", "是否必填"),
        ("IP地址", "设备管理IP，如 192.168.1.1", "必填"),
        ("用户名", "SSH登录用户名", "必填"),
        ("密码",   "SSH登录密码", "必填"),
        ("主机名", "可留空，首次备份成功后会自动回填", "可选"),
        ("端口",   "SSH端口，默认22", "可选"),
        ("", "", ""),
        ("注意", "列名支持中英文，顺序可以随意调整", ""),
        ("注意", "程序只读取数据；仅“主机名”空单元格会被自动回填", ""),
    ]
    for r, row in enumerate(notes, 1):
        for c, val in enumerate(row, 1):
            cell = ws2.cell(row=r, column=c, value=val)
            if r == 1:
                cell.font = Font(bold=True)

    wb.save(output_path)
    print(f"  [✓] Excel 模板已生成: {output_path}")
