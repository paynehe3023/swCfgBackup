# -*- coding: utf-8 -*-
"""
设备配置备份模块
功能：
  - 通过 SSH 登录设备，执行 display current-configuration（只读）
  - 将配置保存为本地 .txt 文件（按日期归档）
  - 自动与上次备份做差异对比（diff），有变更时标注
  - 生成备份汇总报告（Excel）
  - 支持保留策略：每台设备保留最近 N 份（默认 1 份）

【安全说明】
  - 仅执行 display current-configuration（只读命令）
  - 不执行任何写配置命令
  - 不修改设备任何状态
"""

import sys
import os
import time
import re
import json
import difflib
import concurrent.futures
from datetime import datetime
from pathlib import Path             # ← 统一路径
from typing import List, Dict, Optional
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from collector import SSHCollector
from inventory import write_back_hostnames

from utils import fix_console_encoding
fix_console_encoding()

# 保留策略：每台设备保留的备份份数
DEFAULT_KEEP_COUNT = 1
# 失败补采：第一轮结束后等待多久再重试（秒）
DEFAULT_RETRY_DELAY = 30
AUTH_ERROR_MARK = "认证失败"


# ------------------------------------------------------------------ #
#  备份单台设备                                                         #
# ------------------------------------------------------------------ #
def backup_device(ip: str, username: str, password: str,
                  port: int = 22, vendor: str = "auto",
                  backup_dir: str = "backup",
                  cache_dir: str = "", cache_ttl: int = 0) -> Dict:
    """
    登录单台设备，拉取当前配置并保存到本地。
    文件名格式：{sysname}_{ip}_{YYYYMMDD_HHMMSS}.txt
    返回备份结果字典。
    """
    now = datetime.now()
    time_tag = now.strftime("%Y%m%d_%H%M%S")

    result: Dict = {
        "ip": ip, "sysname": "",
        "timestamp": now.strftime("%Y-%m-%d %H:%M:%S"),
        "status": "failed", "filename": "",
        "size_kb": 0, "changed": False,
        "diff_lines": 0, "error": ""
    }

    # 设备子目录：backup/192.168.1.1/  （使用 Path，跨平台兼容）
    dev_dir = Path(backup_dir) / ip
    dev_dir.mkdir(parents=True, exist_ok=True)

    collector = SSHCollector(ip=ip, username=username,
                              password=password, port=port, vendor=vendor)
    if not collector.connect():
        result["error"] = collector.connect_error or "SSH连接失败"
        return result

    try:
        # 自动识别厂商
        if vendor == "auto":
            collector.detect_vendor()

        # 厂商确定后再关一次分页，避免连接初期命令被 banner 吞掉
        collector.disable_paging()

        # ★ 缓存命中检查 ★
        cached_config, cached_vendor, used_cache = _load_config_cache(
            ip, cache_dir, cache_ttl
        )

        if used_cache:
            config_clean = _clean_ansi(cached_config)
            result["status"] = "success (cached)"
            # 与上次备份做差异对比（缓存模式不写新文件）
            changed, diff_lines = _compare_with_latest(dev_dir, None, config_clean)
            file_size = 0
            for f in dev_dir.glob("*.txt"):
                try:
                    file_size = max(file_size, f.stat().st_size / 1024)
                except Exception:
                    pass
            result.update({
                "size_kb": round(file_size, 1),
                "changed": changed,
                "diff_lines": diff_lines,
                "sysname": _extract_sysname(config_clean),
            })
            change_mark = " [有变更!]" if changed else " [无变更]"
            print(f"  → {ip}  [缓存命中]  {change_mark}")
            return result

        # 拉取当前配置（只读命令）；截断时重试一次
        print(f"  → 备份配置: {ip}", end="", flush=True)
        config_clean = ""
        for attempt in (1, 2):
            config_raw = collector.send("display current-configuration",
                                        idle=1.5, max_wait=120)
            candidate = _clean_ansi(config_raw)
            if _config_complete(candidate):
                config_clean = candidate
                break
            if attempt == 1:
                collector.drain()   # 把上一次未输出完的内容读干净再重试
                time.sleep(1)

        if not config_clean:
            result["error"] = "配置不完整（疑似截断），已丢弃本次结果"
            print("  ✗ 配置不完整")
            return result

        result["sysname"] = _extract_sysname(config_clean)

        # 写入配置缓存
        _save_config_cache(ip, cache_dir, config_raw, collector.vendor)

        # ---- 文件名：{sysname}_{ip}_{年月日}_{时分秒}.txt ----
        safe_name = _safe_filename_part(result.get("sysname"), fallback="unknown")
        filename = f"{safe_name}_{ip}_{time_tag}.txt"
        filepath = dev_dir / filename

        # 极端情况（同一秒内重复备份）追加序号
        counter = 0
        while filepath.exists():
            counter += 1
            filepath = dev_dir / f"{safe_name}_{ip}_{time_tag}_{counter}.txt"
            filename = filepath.name

        with open(str(filepath), "w", encoding="utf-8") as f:
            f.write(f"# 设备IP: {ip}\n")
            f.write(f"# 备份时间: {now:%Y-%m-%d %H:%M:%S}\n")
            f.write(f"# 厂商: {collector.vendor}\n")
            f.write("# " + "="*60 + "\n\n")
            f.write(config_clean)

        file_size = filepath.stat().st_size / 1024

        # 差异对比
        changed, diff_lines = _compare_with_latest(dev_dir, filepath, config_clean)

        result.update({
            "status":     "success",
            "filename":   filename,
            "filepath":   str(filepath),
            "size_kb":    round(file_size, 1),
            "changed":    changed,
            "diff_lines": diff_lines,
        })
        change_mark = " [有变更!]" if changed else " [无变更]"
        print(f"  ✓  {file_size:.1f}KB{change_mark}")

    except Exception as e:
        result["error"] = str(e)
        print(f"  ✗  {e}")
    finally:
        collector.disconnect()

    return result


def _clean_ansi(text: str) -> str:
    """去除 ANSI 转义序列、分页提示符和终端控制符"""
    ansi_escape = re.compile(r'\x1B(?:[@-Z\\-_]|\[[0-?]*[ -/]*[@-~])')
    text = ansi_escape.sub("", text)
    text = re.sub(r'-{2,}\s*More\s*-{2,}', '', text, flags=re.I)
    text = text.replace('\b', '')
    text = re.sub(r'\r', '', text)
    return text


# 与配置内容无关、每次备份都会变化的行（用于 diff 与完整性判断）
_RE_ECHO   = re.compile(r"^display current-configuration\s*$", re.I)
_RE_TS     = re.compile(r"^\d{4}-\d{2}-\d{2}[ T]\d{2}:\d{2}:\d{2}(?:\.\d+)?"
                        r"(?:\s*[+-]\d{2}:?\d{2})?\s*$")
_RE_PROMPT = re.compile(r"^[<\[].*[>\]]\s*$")
_RE_SYSNAME = re.compile(r"^\s*sysname\s+(\S+)\s*$", re.M)
_RE_ILLEGAL_FILENAME = re.compile(r'[<>:"/\\|?*\x00-\x1f]')


def _extract_sysname(config: str) -> str:
    """从配置文本中提取设备主机名（sysname）"""
    m = _RE_SYSNAME.search(config or "")
    return m.group(1).strip() if m else ""


def _safe_filename_part(name, fallback: str = "unknown") -> str:
    """把 sysname 转成适合做文件名的片段（去掉 Windows 非法字符）"""
    s = _RE_ILLEGAL_FILENAME.sub("_", str(name or "")).strip(" .")
    s = re.sub(r"\s+", "_", s)
    return s[:64] or fallback


def _is_success(status) -> bool:
    """兼容 'success' 与 'success (cached)'"""
    return str(status or "").startswith("success")


def _config_complete(text: str) -> bool:
    """
    判断 display current-configuration 的输出是否完整。
    华为/H3C 的完整输出以 return 结束，并带设备提示符 <HOST> / [HOST]。
    """
    lines = [l.rstrip() for l in text.splitlines() if l.strip()]
    if len(lines) < 3:          # 空输出或错误提示
        return False
    if not any(l.strip() == "return" for l in lines[-15:]):
        return False
    last = lines[-1].strip()
    return last == "return" or bool(_RE_PROMPT.match(last))


def _normalize_config_lines(text: str) -> list:
    """
    去掉每次都会变化、与配置内容无关的行，供 diff 使用：
      - 文件头元信息（以 # 开头）
      - 设备回显的命令行、设备自己打印的时间戳（仅前几行）
      - 末尾提示符 <HOST> / [HOST]
      - 空行
    """
    lines = text.splitlines()
    while lines and _RE_PROMPT.match(lines[-1].strip()):
        lines.pop()
    out = []
    for line in lines:
        s = line.strip()
        if not s or line.startswith("# "):
            continue
        if len(out) < 3 and (_RE_ECHO.match(s) or _RE_TS.match(s)):
            continue
        out.append(line)
    return out


# ------------------------------------------------------------------ #
#  配置缓存（备份专用）                                               #
# ------------------------------------------------------------------ #
def _config_cache_path(ip: str, cache_dir: str) -> Path:
    return Path(cache_dir) / f"{ip}_config.json"


def _load_config_cache(ip: str, cache_dir: str, cache_ttl: int) -> tuple:
    """
    读取配置缓存。
    返回 (config_text, vendor, is_cached) 三元组。
    cache_ttl <= 0 表示永不过期。
    """
    if not cache_dir:
        return "", "", False
    p = _config_cache_path(ip, cache_dir)
    if not p.exists():
        return "", "", False
    try:
        with open(str(p), "r", encoding="utf-8") as f:
            data = json.load(f)
        if cache_ttl > 0:
            cached_at = datetime.fromisoformat(data.get("_cached_at", "2000-01-01"))
            age_hours = (datetime.now() - cached_at).total_seconds() / 3600
            if age_hours > cache_ttl:
                return "", "", False
        return data.get("config", ""), data.get("vendor", "auto"), True
    except Exception:
        return "", "", False


def _save_config_cache(ip: str, cache_dir: str, config_text: str, vendor: str):
    """将配置文本写入缓存"""
    if not cache_dir:
        return
    try:
        p = _config_cache_path(ip, cache_dir)
        Path(cache_dir).mkdir(parents=True, exist_ok=True)
        with open(str(p), "w", encoding="utf-8") as f:
            json.dump({
                "_cached_at": datetime.now().isoformat(),
                "config": config_text,
                "vendor": vendor,
            }, f, ensure_ascii=False, indent=2)
    except Exception as e:
        print(f"\n  [警告] 缓存写入失败 {ip}: {e}")


def _compare_with_latest(dev_dir: Path, current_file: Path,
                          current_content: str) -> tuple:
    """
    与该设备上一次备份做 diff 对比。
    返回 (是否有变更, 变更行数)
    """
    # 按修改时间取上一份（文件名现在以 sysname 开头，不能再按文件名排序）
    txt_files = sorted([
        f for f in dev_dir.iterdir()
        if f.suffix == ".txt" and f != current_file
    ], key=lambda p: (p.stat().st_mtime_ns, p.name))

    if not txt_files:
        return False, 0   # 首次备份，无对比

    prev_file = txt_files[-1]
    try:
        with open(str(prev_file), "r", encoding="utf-8", errors="replace") as f:
            prev_content = f.read()
    except Exception:
        return False, 0

    # 忽略头部元信息、命令回显、设备时间戳和末尾提示符
    prev_lines = _normalize_config_lines(prev_content)
    curr_lines = _normalize_config_lines(current_content)

    diff = list(difflib.unified_diff(prev_lines, curr_lines, lineterm=""))
    changed_lines = len([l for l in diff if l.startswith(("+", "-"))
                         and not l.startswith(("+++", "---"))])

    if changed_lines > 0 and current_file is not None:
        # diff 文件与配置文件同名，仅扩展名不同
        diff_path = current_file.with_suffix(".diff")
        with open(str(diff_path), "w", encoding="utf-8") as f:
            f.write(f"# 差异对比：{prev_file.name} → {current_file.name}\n")
            f.write(f"# 变更行数: {changed_lines}\n\n")
            f.write("\n".join(diff))

    return changed_lines > 0, changed_lines


# ------------------------------------------------------------------ #
#  清理旧备份（保留策略）                                               #
# ------------------------------------------------------------------ #
def cleanup_old_backups(backup_dir: str,
                        keep_count: int = DEFAULT_KEEP_COUNT):
    """
    按份数清理旧备份：
      - 每台设备只保留最近 keep_count 份 .txt（及其同名 .diff）
      - 更早的 .txt / .diff 全部删除，不再按天数判断
      - 没有对应 .txt 的 .diff 一并删除
      - 仅删除 .txt 和 .diff，不删除目录本身；以 _ 开头的目录（如 _truncated）跳过
    """
    keep_count = max(1, int(keep_count))
    backup_path = Path(backup_dir)
    deleted = 0
    if not backup_path.is_dir():
        return
    for dev_dir in backup_path.iterdir():
        if not dev_dir.is_dir() or dev_dir.name.startswith("_"):
            continue
        try:
            txt_files = sorted(
                [p for p in dev_dir.iterdir() if p.suffix == ".txt"],
                key=lambda p: p.stat().st_mtime,
                reverse=True,
            )
        except Exception:
            continue
        keep = {p.stem for p in txt_files[:keep_count]}
        for fpath in txt_files:
            if fpath.stem in keep:
                continue
            try:
                fpath.unlink()
                deleted += 1
            except Exception:
                pass   # 删除失败不中断流程
        # 只保留与保留配置同名的 .diff
        for fpath in list(dev_dir.iterdir()):
            if fpath.suffix != ".diff" or fpath.stem in keep:
                continue
            try:
                fpath.unlink()
                deleted += 1
            except Exception:
                pass
    if deleted:
        print(f"  [清理] 已删除 {deleted} 个旧备份文件"
              f"（每设备保留最近 {keep_count} 份）")


# ------------------------------------------------------------------ #
#  生成备份汇总报告                                                     #
# ------------------------------------------------------------------ #
def generate_backup_report(results: List[Dict], output_path: str):
    """生成本次备份汇总 Excel 报告"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "备份汇总"

    # 统计
    total   = len(results)
    ok      = sum(1 for r in results if _is_success(r["status"]))
    changed = sum(1 for r in results if r.get("changed"))
    failed  = total - ok

    # 标题
    ws.merge_cells("A1:I1")
    ws["A1"] = f"配置备份报告  |  {datetime.now():%Y-%m-%d %H:%M}  |  总计:{total}  成功:{ok}  变更:{changed}  失败:{failed}"
    ws["A1"].font = Font(bold=True, size=12, color="1F4E79")
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 25

    headers = ["IP地址", "主机名", "备份时间", "状态", "文件大小(KB)",
               "配置变更", "变更行数", "文件名", "错误信息"]
    fill_h = PatternFill("solid", fgColor="1F4E79")
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=2, column=col, value=h)
        c.fill = fill_h
        c.font = Font(bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="center")

    fill_ok      = PatternFill("solid", fgColor="E2EFDA")
    fill_fail    = PatternFill("solid", fgColor="FFE7E7")
    fill_changed = PatternFill("solid", fgColor="FFF2CC")

    for row_idx, r in enumerate(results, 3):
        if not _is_success(r["status"]):
            bg = fill_fail
        elif r.get("changed"):
            bg = fill_changed
        else:
            bg = fill_ok

        status_text = "✓ 成功" if _is_success(r["status"]) else f"✗ 失败"
        changed_text = "⚠ 有变更" if r.get("changed") else ("✓ 无变更" if _is_success(r["status"]) else "-")

        row_data = [
            r["ip"],
            r.get("sysname", ""),
            r["timestamp"],
            status_text,
            r.get("size_kb", "-"),
            changed_text,
            r.get("diff_lines", "-"),
            r.get("filename", "-"),
            r.get("error", ""),
        ]
        for col, val in enumerate(row_data, 1):
            c = ws.cell(row=row_idx, column=col, value=val)
            c.fill = bg
            c.alignment = Alignment(vertical="center")

    # 列宽
    widths = [16, 22, 18, 10, 14, 12, 10, 38, 30]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(2, col).column_letter].width = w

    wb.save(output_path)
    print(f"  [✓] 备份报告已保存: {output_path}")


# ------------------------------------------------------------------ #
#  批量备份主函数                                                       #
# ------------------------------------------------------------------ #
def _run_pass(devices: List[Dict], backup_path: Path, concurrency: int,
              cache_dir: str, cache_ttl: int) -> List[Dict]:
    """跑一轮备份（并发），返回结果列表。"""
    results = []
    if not devices:
        return results
    with concurrent.futures.ThreadPoolExecutor(max_workers=concurrency) as executor:
        futures = {
            executor.submit(
                backup_device,
                dev["ip"], dev["username"], dev["password"],
                dev.get("port", 22), dev.get("vendor", "auto"),
                str(backup_path), cache_dir, cache_ttl
            ): dev
            for dev in devices
        }
        for future in concurrent.futures.as_completed(futures):
            try:
                results.append(future.result())
            except Exception as e:
                dev = futures[future]
                results.append({"ip": dev["ip"], "status": "failed",
                                "error": str(e), "changed": False,
                                "diff_lines": 0, "size_kb": 0,
                                "filename": "", "timestamp": ""})
    return results


def select_retry_devices(devices: List[Dict], results: List[Dict],
                         retry_auth: bool = False) -> List[Dict]:
    """挑出需要补采的失败设备；默认跳过认证失败，避免连续错误导致锁号。"""
    by_ip = {r["ip"]: r for r in results}
    retry = []
    for dev in devices:
        r = by_ip.get(dev["ip"])
        if r is None or _is_success(r.get("status")):
            continue
        if not retry_auth and AUTH_ERROR_MARK in str(r.get("error", "")):
            continue
        if "算法不兼容" in str(r.get("error", "")):
            continue   # 客户端/设备算法不匹配，重试无意义
        retry.append(dev)
    return retry


def merge_results(results: List[Dict], retry_results: List[Dict]) -> List[Dict]:
    """用补采结果覆盖同一 IP 的旧结果，并按 IP 排序。"""
    by_ip = {r["ip"]: r for r in results}
    by_ip.update({r["ip"]: r for r in retry_results})
    return sorted(by_ip.values(), key=lambda r: r["ip"])


def run_backup(devices: List[Dict], backup_dir: str = "backup",
               concurrency: int = 3,
               cache_dir: str = "", cache_ttl: int = 0,
               keep_count: int = DEFAULT_KEEP_COUNT,
               retry_failed: bool = True,
               retry_delay: int = DEFAULT_RETRY_DELAY,
               retry_auth: bool = False,
               inventory_path: str = "", sheet_name: str = None,
               write_back: bool = True) -> List[Dict]:
    """
    批量备份所有设备配置。
    concurrency 建议 3~5（备份比采集耗时，配置文件可能很大）
    retry_failed: 第一轮结束后对失败设备补采一次
    retry_auth:   认证失败是否也补采（默认否，避免锁号）
    """
    backup_path = Path(backup_dir)
    backup_path.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

    print(f"\n{'='*55}")
    print(f"  配置备份任务启动")
    print(f"  设备总数: {len(devices)}  并发: {concurrency}")
    print(f"  备份目录: {backup_path.resolve()}")
    print(f"  保留策略: 每设备保留最近 {keep_count} 份")
    print(f"{'='*55}\n")

    results = _run_pass(devices, backup_path, concurrency, cache_dir, cache_ttl)

    # ---------- 失败补采 ----------
    if retry_failed:
        retry_devices = select_retry_devices(devices, results, retry_auth)
        if retry_devices:
            print(f"\n  [补采] {len(retry_devices)} 台失败设备将在 "
                  f"{retry_delay} 秒后重试一次 ...\n")
            time.sleep(max(0, retry_delay))
            retry_results = _run_pass(retry_devices, backup_path,
                                      concurrency, cache_dir, cache_ttl)
            results = merge_results(results, retry_results)

    results.sort(key=lambda r: r["ip"])

    # 清理旧备份
    cleanup_old_backups(str(backup_path), keep_count)

    # 生成报告
    report_path = backup_path / f"backup_report_{timestamp}.xlsx"
    generate_backup_report(results, str(report_path))

    # 主机名回填：只填空单元格，不覆盖手工维护的值
    if write_back and inventory_path:
        hostmap = {r["ip"]: r.get("sysname", "")
                   for r in results if _is_success(r["status"])}
        hostmap = {ip: name for ip, name in hostmap.items() if name}
        if hostmap:
            info = write_back_hostnames(inventory_path, hostmap, sheet_name)
            if info["error"]:
                print(f"  [主机名] 回填失败: {info['error']}")
            else:
                print(f"  [主机名] 已回填 {info['filled']} 台到 "
                      f"{inventory_path}（保留已有值 {info['skipped']} 台）")

    ok      = sum(1 for r in results if _is_success(r["status"]))
    changed = sum(1 for r in results if r.get("changed"))
    print(f"\n  备份完成！成功: {ok}/{len(results)}  有配置变更: {changed} 台")
    print(f"{'='*55}\n")

    return results
